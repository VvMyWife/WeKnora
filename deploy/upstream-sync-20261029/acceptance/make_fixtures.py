from pathlib import Path
import io,struct,zipfile,json
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
out=Path('/fixtures')
# Raster-only PDF, with no embedded text layer.
im=Image.new('RGB',(1200,1600),'white');draw=ImageDraw.Draw(im)
draw.text((90,140),'LJC SCAN ACCEPTANCE 20261029',fill='black',font_size=36)
draw.text((90,250),'Archive boxes: 17',fill='black',font_size=48)
im.save(out/'scan.png')
im.save(out/'scan.pdf',format='PDF',resolution=144)
# A synthetic, valid placeable WMF with window mapping and a rectangle.
def rec(fn,*params):return struct.pack('<IH',3+len(params),fn)+struct.pack('<'+'H'*len(params),*params)
records=rec(0x0103,8)+rec(0x020B,0,0)+rec(0x020C,1000,1000)+rec(0x041B,900,900,100,100)+rec(0)
header=struct.pack('<I H hhhh H I',0x9AC6CDD7,0,0,0,1000,1000,144,0)
checksum=0
for word in struct.unpack('<10H',header):checksum^=word
wmf=header+struct.pack('<H',checksum)+struct.pack('<HHHIHIH',1,9,0x300,(18+len(records))//2,0,7,0)+records
(out/'rectangle.wmf').write_bytes(wmf)
deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);pic=slide.shapes.add_picture(str(out/'scan.png'),Inches(1),Inches(1),width=Inches(3));pic.name='ProofVector';b=io.BytesIO();deck.save(b)
with zipfile.ZipFile(io.BytesIO(b.getvalue())) as z:files={n:z.read(n) for n in z.namelist()}
media=next(n for n in files if n.startswith('ppt/media/'));vector=media.rsplit('.',1)[0]+'.wmf'
files[vector]=wmf;del files[media]
for n in list(files):
 if n.endswith('.rels'):files[n]=files[n].replace(media.split('/')[-1].encode(),vector.split('/')[-1].encode())
files['[Content_Types].xml']=files['[Content_Types].xml'].replace(b'Extension="png" ContentType="image/png"',b'Extension="wmf" ContentType="image/x-wmf"')
with zipfile.ZipFile(out/'vector.pptx','w',zipfile.ZIP_DEFLATED) as z:
 for n,data in files.items():z.writestr(n,data)
wb=Workbook();ws=wb.active;ws.append(['Item','Quantity']);ws.append(['Archive boxes',17]);ws.append(['N/A',0]);wb.save(out/'table.xlsx')
(out/'table.csv').write_text('Item,Quantity\nArchive boxes,17\nN/A,0\n')
import sys; sys.path.insert(0,"/app")
from docreader.parser.pptx_media import rasterize_media_bytes
raster=rasterize_media_bytes('rectangle.wmf',wmf)
assert raster,'valid WMF did not rasterize'
(out/'wmf-raster.png').write_bytes(raster)
print(json.dumps({'scan_pdf_text_layer':b'/Font' in (out/'scan.pdf').read_bytes(),'wmf_raster_bytes':len(raster),'fixtures':[p.name for p in out.iterdir()]}))
